#!/usr/bin/env python3
"""
Baseball Ticket Extractor V2

Extracts structured data from baseball ticket images using Claude's vision API,
validates the response, performs deterministic cross-field checks, and writes
an Excel workbook designed for archival/review workflows.

Requirements:
    pip install anthropic openpyxl pillow pydantic pillow-heif

Usage:
    python extract_tickets_v2.py --tickets ./tickets --output tickets.xlsx
    python extract_tickets_v2.py --tickets ./tickets --output tickets.xlsx --model sonnet
    python extract_tickets_v2.py --tickets ./tickets --output tickets.xlsx --workers 6
"""

from __future__ import annotations

import argparse
import base64
import io
import json
import logging
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any, Literal

import anthropic
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps, ImageEnhance
from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator


# ─────────────────────────────────────────────────────────────────────────────
# Configuration
# ─────────────────────────────────────────────────────────────────────────────

SUPPORTED_EXTENSIONS = {
    ".jpg", ".jpeg", ".png", ".heic", ".webp", ".bmp", ".tif", ".tiff"
}

MODELS = {
    "haiku": "claude-haiku-4-5-20251001",
    "sonnet": "claude-sonnet-4-6",
}

DEFAULT_MAX_IMAGE_DIMENSION = 2400
DEFAULT_JPEG_QUALITY = 90
DEFAULT_WORKERS = 4
DEFAULT_RETRIES = 3
DEFAULT_TIMEOUT = 120.0

ALLOWED_CONFIDENCE = {"high", "medium", "low"}
ALLOWED_SEASON_TICKET = {"Yes", "No"}

EXTRACTION_PROMPT = """
You are extracting information from ONE baseball ticket image.

Your job is transcription/extraction, not reconstruction or inference.

CRITICAL RULES:
1. Extract only information supported by text or clearly visible printed content.
2. NEVER guess an obscured, ambiguous, or unreadable value.
3. NEVER infer a value from baseball knowledge, team schedules, stadium history,
   likely opponents, or what you think the ticket "should" say.
4. If a field cannot be read confidently from the image, use null.
5. Preserve the ticket's wording where practical. Do not silently "correct"
   unusual spellings or historical venue/team names.
6. For raw_text, provide a concise transcription of the legible ticket text,
   preserving line breaks approximately. Do not invent missing text.
7. For evidence fields, quote only short, visibly legible text from the ticket.
8. A field's confidence describes your confidence that the visible text supports
   THAT FIELD, not your confidence based on outside knowledge.
9. The overall confidence should be the lowest confidence among important
   extracted fields.
10. If date is visible but day-of-week is not printed, set day_of_week_printed
    to null. Python will calculate a day-of-week independently from the date.
11. Do not calculate, derive, or repair the game number, seat, row, section,
    gate, price, or teams from context.
12. Return ONLY valid JSON matching the requested schema.

Return this exact top-level structure:

{
  "raw_text": "string or null",
  "date": {"value": "string or null", "confidence": "high|medium|low", "evidence": "string or null"},
  "day_of_week_printed": {"value": "string or null", "confidence": "high|medium|low", "evidence": "string or null"},
  "game_time": {"value": "string or null", "confidence": "high|medium|low", "evidence": "string or null"},
  "home_team": {"value": "string or null", "confidence": "high|medium|low", "evidence": "string or null"},
  "away_team": {"value": "string or null", "confidence": "high|medium|low", "evidence": "string or null"},
  "stadium": {"value": "string or null", "confidence": "high|medium|low", "evidence": "string or null"},
  "game_number": {"value": "string or null", "confidence": "high|medium|low", "evidence": "string or null"},
  "section": {"value": "string or null", "confidence": "high|medium|low", "evidence": "string or null"},
  "row": {"value": "string or null", "confidence": "high|medium|low", "evidence": "string or null"},
  "seat": {"value": "string or null", "confidence": "high|medium|low", "evidence": "string or null"},
  "seat_type": {"value": "string or null", "confidence": "high|medium|low", "evidence": "string or null"},
  "gate": {"value": "string or null", "confidence": "high|medium|low", "evidence": "string or null"},
  "price": {"value": "string or null", "confidence": "high|medium|low", "evidence": "string or null"},
  "season_ticket": {"value": "Yes|No|null", "confidence": "high|medium|low", "evidence": "string or null"},
  "notes": "string or null",
  "confidence": "high|medium|low"
}
""".strip()


# ─────────────────────────────────────────────────────────────────────────────
# Logging
# ─────────────────────────────────────────────────────────────────────────────

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
)
logger = logging.getLogger("ticket-extractor")


# ─────────────────────────────────────────────────────────────────────────────
# Pydantic schema
# ─────────────────────────────────────────────────────────────────────────────

class FieldExtraction(BaseModel):
    model_config = ConfigDict(extra="forbid")

    value: str | None = None
    confidence: Literal["high", "medium", "low"]
    evidence: str | None = None


class TicketExtraction(BaseModel):
    model_config = ConfigDict(extra="forbid")

    raw_text: str | None = None
    date: FieldExtraction
    day_of_week_printed: FieldExtraction
    game_time: FieldExtraction
    home_team: FieldExtraction
    away_team: FieldExtraction
    stadium: FieldExtraction
    game_number: FieldExtraction
    section: FieldExtraction
    row: FieldExtraction
    seat: FieldExtraction
    seat_type: FieldExtraction
    gate: FieldExtraction
    price: FieldExtraction
    season_ticket: FieldExtraction
    notes: str | None = None
    confidence: Literal["high", "medium", "low"]

    @field_validator("season_ticket")
    @classmethod
    def validate_season_ticket(cls, value: FieldExtraction) -> FieldExtraction:
        if value.value is not None and value.value not in ALLOWED_SEASON_TICKET:
            raise ValueError("season_ticket.value must be Yes, No, or null")
        return value


# ─────────────────────────────────────────────────────────────────────────────
# Image handling
# ─────────────────────────────────────────────────────────────────────────────

def normalize_image(path: Path, max_dimension: int = DEFAULT_MAX_IMAGE_DIMENSION) -> tuple[str, str]:
    """
    Normalize orientation and size, then return base64 JPEG data.

    JPEG is used for all supported formats so the API receives a consistent
    media type. Original files are never modified.
    """
    try:
        if path.suffix.lower() == ".heic":
            try:
                from pillow_heif import register_heif_opener
                register_heif_opener()
            except ImportError as exc:
                raise RuntimeError(
                    "HEIC image encountered but pillow-heif is not installed. "
                    "Install with: pip install pillow-heif"
                ) from exc

        with Image.open(path) as image:
            image = ImageOps.exif_transpose(image)
            image = image.convert("RGB")

            width, height = image.size
            largest = max(width, height)
            if largest > max_dimension:
                scale = max_dimension / largest
                image = image.resize(
                    (max(1, round(width * scale)), max(1, round(height * scale))),
                    Image.Resampling.LANCZOS,
                )

            # Mild enhancement; deliberately conservative to avoid altering text.
            image = ImageEnhance.Contrast(image).enhance(1.05)
            image = ImageEnhance.Sharpness(image).enhance(1.10)

            buf = io.BytesIO()
            image.save(buf, format="JPEG", quality=DEFAULT_JPEG_QUALITY, optimize=True)
            return base64.standard_b64encode(buf.getvalue()).decode("ascii"), "image/jpeg"

    except Exception as exc:
        raise RuntimeError(f"Could not normalize image '{path}': {exc}") from exc


# ─────────────────────────────────────────────────────────────────────────────
# API / extraction
# ─────────────────────────────────────────────────────────────────────────────

def _strip_json_fences(text: str) -> str:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
        text = re.sub(r"\s*```$", "", text)
    return text.strip()


def _is_retryable(exc: Exception) -> bool:
    retryable_types = (
        anthropic.RateLimitError,
        anthropic.APITimeoutError,
        anthropic.APIConnectionError,
        anthropic.InternalServerError,
    )
    return isinstance(exc, retryable_types)


def extract_ticket(
    client: anthropic.Anthropic,
    image_path: Path,
    model: str,
    retries: int = DEFAULT_RETRIES,
    timeout: float = DEFAULT_TIMEOUT,
) -> TicketExtraction:
    """Extract and validate one ticket, retrying transient API failures."""
    image_data, media_type = normalize_image(image_path)

    last_error: Exception | None = None

    for attempt in range(1, retries + 1):
        try:
            response = client.messages.create(
                model=model,
                max_tokens=2500,
                timeout=timeout,
                messages=[{
                    "role": "user",
                    "content": [
                        {
                            "type": "image",
                            "source": {
                                "type": "base64",
                                "media_type": media_type,
                                "data": image_data,
                            },
                        },
                        {"type": "text", "text": EXTRACTION_PROMPT},
                    ],
                }],
            )

            raw = "".join(
                block.text for block in response.content
                if getattr(block, "type", None) == "text"
            ).strip()

            if not raw:
                raise ValueError("Claude returned no text content")

            payload = json.loads(_strip_json_fences(raw))
            return TicketExtraction.model_validate(payload)

        except (json.JSONDecodeError, ValidationError, ValueError) as exc:
            # These are deterministic output problems; retrying can sometimes
            # help, but they are classified separately from transport failures.
            last_error = exc
            if attempt < retries:
                time.sleep(1.5 * attempt)
                continue
            raise RuntimeError(f"Invalid model output after {retries} attempts: {exc}") from exc

        except Exception as exc:
            last_error = exc
            if not _is_retryable(exc) or attempt >= retries:
                raise
            delay = min(30.0, 2 ** (attempt - 1))
            logger.warning(
                "Transient API error for %s (attempt %d/%d): %s; retrying in %.1fs",
                image_path.name, attempt, retries, exc, delay,
            )
            time.sleep(delay)

    raise RuntimeError(f"Extraction failed: {last_error}")


# ─────────────────────────────────────────────────────────────────────────────
# Deterministic validation
# ─────────────────────────────────────────────────────────────────────────────

_LEADING_DOW_RE = re.compile(
    r"^(mon|tue|wed|thu|fri|sat|sun)[a-z]*\.?,?\s+", re.IGNORECASE
)


def _parse_date(value: str | None) -> datetime | None:
    if not value:
        return None

    cleaned = value.strip()
    cleaned = _LEADING_DOW_RE.sub("", cleaned)  # drop "TUE " / "Tue., " prefix
    cleaned = cleaned.replace(".", "")          # "JUL." -> "JUL", "SEPT." -> "SEPT"
    cleaned = re.sub(r"\s+", " ", cleaned).strip()

    formats = (
        "%B %d, %Y", "%B %d %Y",
        "%b %d, %Y", "%b %d %Y",
        "%m/%d/%Y", "%m/%d/%y",
        "%Y-%m-%d",
    )
    for fmt in formats:
        try:
            return datetime.strptime(cleaned, fmt)
        except ValueError:
            continue

    # Last resort: normalize 4-letter month abbreviations (e.g. "Sept") to 3.
    cleaned2 = re.sub(r"\bsept\b", "sep", cleaned, flags=re.IGNORECASE)
    if cleaned2 != cleaned:
        for fmt in ("%b %d, %Y", "%b %d %Y"):
            try:
                return datetime.strptime(cleaned2, fmt)
            except ValueError:
                continue

    return None


def _normalize_day(value: str) -> str:
    """Reduce a weekday string to a bare 3-letter lowercase code so that
    'Wednesday', 'WED', 'Wed.', and 'wed' all compare equal."""
    return re.sub(r"[^A-Za-z]", "", value).lower()[:3]


def validate_extraction(data: TicketExtraction) -> list[str]:
    """
    Deterministic checks. These do not 'correct' Claude; they flag records
    that deserve human review.
    """
    issues: list[str] = []

    date_value = data.date.value
    parsed_date = _parse_date(date_value)

    if date_value and parsed_date is None:
        issues.append(f"Date could not be parsed: {date_value!r}")

    if parsed_date and data.day_of_week_printed.value:
        calculated = parsed_date.strftime("%A")
        printed = data.day_of_week_printed.value.strip()
        if _normalize_day(calculated) != _normalize_day(printed):
            issues.append(
                f"Day mismatch: ticket says {printed!r}, date calculates to {calculated!r}"
            )

    if (
        data.home_team.value
        and data.away_team.value
        and data.home_team.value.strip().lower() == data.away_team.value.strip().lower()
    ):
        issues.append("Home team and away team are identical")

    if data.price.value and not re.search(r"\d", data.price.value):
        issues.append(f"Price contains no digits: {data.price.value!r}")

    for field_name in (
        "game_number", "section", "row", "seat", "gate"
    ):
        field = getattr(data, field_name)
        if field.value is not None and not field.value.strip():
            issues.append(f"{field_name} is blank instead of null")

    if data.confidence == "high":
        low_or_medium = []
        for field_name in (
            "date", "game_time", "home_team", "away_team",
            "stadium", "section", "row", "seat"
        ):
            field = getattr(data, field_name)
            if field.value is not None and field.confidence != "high":
                low_or_medium.append(field_name)
        if low_or_medium:
            issues.append(
                "Overall confidence is high but these populated important fields "
                f"are not high confidence: {', '.join(low_or_medium)}"
            )

    return issues


# ─────────────────────────────────────────────────────────────────────────────
# Excel output
# ─────────────────────────────────────────────────────────────────────────────

HEADERS = [
    "File Name", "Date", "Day Printed", "Day Calculated", "Time",
    "Home Team", "Away Team", "Stadium", "Game #", "Section", "Row", "Seat",
    "Seat Type", "Gate", "Price", "Season Ticket", "Confidence",
    "Validation Issues", "Notes", "Raw Text", "Model", "Processed At",
    "Status", "Error",
]

COLUMN_WIDTHS = {
    "File Name": 30, "Date": 18, "Day Printed": 14, "Day Calculated": 14,
    "Time": 12, "Home Team": 22, "Away Team": 22, "Stadium": 24,
    "Game #": 9, "Section": 10, "Row": 8, "Seat": 8, "Seat Type": 20,
    "Gate": 8, "Price": 12, "Season Ticket": 14, "Confidence": 12,
    "Validation Issues": 45, "Notes": 40, "Raw Text": 55, "Model": 30,
    "Processed At": 22, "Status": 14, "Error": 45,
}

CONFIDENCE_COLORS = {
    "high": "C6EFCE",
    "medium": "FFEB9C",
    "low": "FFC7CE",
    "error": "FFC7CE",
    "review": "FFEB9C",
}

HEADER_FILL = PatternFill("solid", start_color="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
DATA_FONT = Font(name="Arial", size=10)
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _field_value(data: TicketExtraction, name: str) -> str | None:
    return getattr(data, name).value


def build_excel(records: list[dict[str, Any]], output_path: Path, model: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tickets"
    ws.freeze_panes = "B2"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 32

    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS.get(header, 14)

    for row_idx, rec in enumerate(records, 2):
        data: TicketExtraction | None = rec.get("data")
        issues = rec.get("validation_issues", [])
        error = rec.get("error")
        processed_at = rec.get("processed_at")

        if data:
            parsed_date = _parse_date(data.date.value)
            calculated_day = parsed_date.strftime("%A") if parsed_date else None

            values = [
                rec["filename"],
                _field_value(data, "date"),
                _field_value(data, "day_of_week_printed"),
                calculated_day,
                _field_value(data, "game_time"),
                _field_value(data, "home_team"),
                _field_value(data, "away_team"),
                _field_value(data, "stadium"),
                _field_value(data, "game_number"),
                _field_value(data, "section"),
                _field_value(data, "row"),
                _field_value(data, "seat"),
                _field_value(data, "seat_type"),
                _field_value(data, "gate"),
                _field_value(data, "price"),
                _field_value(data, "season_ticket"),
                data.confidence,
                "\n".join(issues) if issues else None,
                data.notes,
                data.raw_text,
                model,
                processed_at,
                "OK" if not issues else "REVIEW",
                None,
            ]
            row_status = "review" if issues else data.confidence
        else:
            values = [
                rec["filename"], *([None] * 20),
                model, processed_at, "ERROR", error
            ]
            row_status = "error"

        ws.row_dimensions[row_idx].height = 60 if data else 35

        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = DATA_FONT
            cell.border = BORDER
            cell.alignment = CENTER if col_idx in {
                1, 2, 3, 4, 5, 9, 10, 11, 12, 14, 15, 16, 17, 23
            } else LEFT

        fill = PatternFill("solid", start_color=CONFIDENCE_COLORS.get(row_status, "FFFFFF"))
        for col_idx in range(1, len(HEADERS) + 1):
            ws.cell(row=row_idx, column=col_idx).fill = fill

    if records:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(records) + 1}"

    ws.conditional_formatting.add(
        f"Q2:Q{max(2, len(records) + 1)}",
        FormulaRule(formula=['Q2="high"'], fill=PatternFill("solid", start_color=CONFIDENCE_COLORS["high"]))
    )

    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2.sheet_view.showGridLines = False
    total = len(records)
    succeeded = sum(1 for r in records if r.get("data"))
    failed = total - succeeded
    review = sum(1 for r in records if r.get("validation_issues"))
    high = sum(
        1 for r in records
        if r.get("data") and r["data"].confidence == "high"
    )
    medium = sum(
        1 for r in records
        if r.get("data") and r["data"].confidence == "medium"
    )
    low = sum(
        1 for r in records
        if r.get("data") and r["data"].confidence == "low"
    )

    ws2["A1"] = "Baseball Ticket Extraction V2"
    ws2["A1"].font = Font(bold=True, size=15, name="Arial", color="1F4E79")
    ws2.column_dimensions["A"].width = 38
    ws2.column_dimensions["B"].width = 20

    summary_rows = [
        ("Total tickets processed", total),
        ("Successfully extracted", succeeded),
        ("API / processing errors", failed),
        ("Records requiring review", review),
        ("High confidence", high),
        ("Medium confidence", medium),
        ("Low confidence", low),
        ("Model", model),
        ("Run date", datetime.now().strftime("%B %d, %Y %I:%M %p")),
    ]

    for i, (label, value) in enumerate(summary_rows, 3):
        ws2.cell(row=i, column=1, value=label).font = Font(bold=True, name="Arial")
        ws2.cell(row=i, column=2, value=value).font = Font(name="Arial")

    ws2["A14"] = "Review guidance"
    ws2["A14"].font = Font(bold=True, name="Arial", color="1F4E79")
    ws2["A15"] = (
        "Review rows marked REVIEW or ERROR. Validation flags are deterministic "
        "checks and do not automatically mean the extraction is wrong."
    )
    ws2["A15"].alignment = LEFT
    ws2.merge_cells("A15:B17")

    wb.save(output_path)


# ─────────────────────────────────────────────────────────────────────────────
# CLI / batch processing
# ─────────────────────────────────────────────────────────────────────────────

def process_one(
    client: anthropic.Anthropic,
    image_path: Path,
    model: str,
    retries: int,
    timeout: float,
) -> dict[str, Any]:
    processed_at = datetime.now().isoformat(timespec="seconds")

    try:
        data = extract_ticket(
            client=client,
            image_path=image_path,
            model=model,
            retries=retries,
            timeout=timeout,
        )
        issues = validate_extraction(data)
        return {
            "filename": image_path.name,
            "data": data,
            "validation_issues": issues,
            "processed_at": processed_at,
        }
    except Exception as exc:
        return {
            "filename": image_path.name,
            "data": None,
            "validation_issues": [],
            "processed_at": processed_at,
            "error": f"{type(exc).__name__}: {exc}",
        }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Extract baseball ticket data to a review-friendly Excel workbook"
    )
    parser.add_argument("--tickets", required=True, help="Folder containing ticket images")
    parser.add_argument(
        "--output", default="baseball_tickets_v2.xlsx",
        help="Output Excel path"
    )
    parser.add_argument(
        "--model", choices=sorted(MODELS), default="haiku",
        help="Claude model alias (default: haiku)"
    )
    parser.add_argument(
        "--api-key", default=None,
        help="Anthropic API key (or use ANTHROPIC_API_KEY)"
    )
    parser.add_argument(
        "--workers", type=int, default=DEFAULT_WORKERS,
        help=f"Maximum concurrent API requests (default: {DEFAULT_WORKERS})"
    )
    parser.add_argument(
        "--retries", type=int, default=DEFAULT_RETRIES,
        help=f"Retries for transient/invalid responses (default: {DEFAULT_RETRIES})"
    )
    parser.add_argument(
        "--timeout", type=float, default=DEFAULT_TIMEOUT,
        help=f"Per-request timeout in seconds (default: {DEFAULT_TIMEOUT})"
    )
    parser.add_argument(
        "--max-image-dimension", type=int, default=DEFAULT_MAX_IMAGE_DIMENSION,
        help=f"Resize largest image dimension to this many pixels (default: {DEFAULT_MAX_IMAGE_DIMENSION})"
    )
    args = parser.parse_args()

    if args.workers < 1:
        parser.error("--workers must be >= 1")
    if args.retries < 1:
        parser.error("--retries must be >= 1")

    tickets_dir = Path(args.tickets)
    if not tickets_dir.is_dir():
        print(f"Error: '{tickets_dir}' is not a valid directory.", file=sys.stderr)
        return 1

    image_files = sorted(
        p for p in tickets_dir.iterdir()
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS
    )

    if not image_files:
        print(f"No supported image files found in '{tickets_dir}'.", file=sys.stderr)
        return 1

    model_id = MODELS[args.model]
    client = anthropic.Anthropic(
        api_key=args.api_key or os.getenv("ANTHROPIC_API_KEY")
    )

    print("\n🎟️  Baseball Ticket Extractor V2")
    print(f"   Found    : {len(image_files)} image(s)")
    print(f"   Model    : {args.model} ({model_id})")
    print(f"   Workers  : {args.workers}")
    print(f"   Output   : {args.output}\n")

    records: list[dict[str, Any]] = [None] * len(image_files)  # type: ignore[list-item]

    with ThreadPoolExecutor(max_workers=args.workers) as executor:
        future_map = {
            executor.submit(
                process_one, client, path, model_id, args.retries, args.timeout
            ): (index, path)
            for index, path in enumerate(image_files)
        }

        completed = 0
        for future in as_completed(future_map):
            index, path = future_map[future]
            record = future.result()
            records[index] = record
            completed += 1

            if record.get("error"):
                print(
                    f"[{completed:>3}/{len(image_files)}] {path.name} ... ✗ "
                    f"{record['error']}"
                )
            else:
                data: TicketExtraction = record["data"]
                status = "REVIEW" if record["validation_issues"] else data.confidence.upper()
                print(f"[{completed:>3}/{len(image_files)}] {path.name} ... ✓ ({status})")

    output_path = Path(args.output)
    build_excel(records, output_path, model_id)

    succeeded = sum(1 for r in records if r.get("data"))
    errors = len(records) - succeeded
    reviews = sum(1 for r in records if r.get("validation_issues"))

    print(f"\n✅ Done! {succeeded}/{len(records)} tickets extracted successfully.")
    print(f"   Review flags : {reviews}")
    print(f"   Errors       : {errors}")
    print(f"   Saved to     : {output_path.resolve()}\n")

    return 0 if errors == 0 else 2


if __name__ == "__main__":
    raise SystemExit(main())
